import io
import hashlib
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app as fastapi_app
from app.core.security import create_access_token
from app.core.database import SessionLocal
from app.core.auth import get_current_user
from app.models.user import User
from app.models.assistant import AssistantConversation, AssistantMessage
from app.models.document import Document
from app.ai.providers.mock_provider import MockAIProvider
from app.ai.base import BaseAIProvider

client = TestClient(fastapi_app)


def get_auth_headers(email: str = "assistant_test_user@litebooks.ai") -> dict:
    db = SessionLocal()
    try:
        user = db.query(User).filter(User.email == email).first()
        if not user:
            user = User(
                username=email.split("@")[0],
                email=email,
                full_name="Assistant Test User",
                hashed_password="hashed_pw_test",
                role="ACCOUNTANT",
                is_active=True,
            )
            db.add(user)
            db.commit()
            db.refresh(user)
        else:
            user.is_active = True
            db.commit()
        token = create_access_token(data={"sub": str(user.id)})
        return {"Authorization": f"Bearer {token}"}
    finally:
        db.close()


def get_other_user_headers(email: str = "other_user@litebooks.ai") -> dict:
    db = SessionLocal()
    try:
        user = db.query(User).filter(User.email == email).first()
        if not user:
            user = User(
                username=email.split("@")[0],
                email=email,
                full_name="Other User",
                hashed_password="hashed_pw_test",
                role="ACCOUNTANT",
                is_active=True,
            )
            db.add(user)
            db.commit()
            db.refresh(user)
        else:
            user.is_active = True
            db.commit()
        token = create_access_token(data={"sub": str(user.id)})
        return {"Authorization": f"Bearer {token}"}
    finally:
        db.close()


def create_sample_financial_workbook_bytes() -> bytes:
    wb = Workbook()
    
    # Balance Sheet
    ws_bs = wb.active
    ws_bs.title = "Balance Sheet"
    ws_bs.append(["Zion Lifts LLP - Balance Sheet as at 31st March 2026"])
    ws_bs.append(["Particulars", "Note No.", "Figures as at 31-03-2026", "Figures as at 31-03-2025"])
    ws_bs.append(["I. EQUITY AND LIABILITIES", "", "", ""])
    ws_bs.append(["Partner's Funds", "", "", ""])
    ws_bs.append(["Capital Account", "1", 10000000.00, 9000000.00])
    ws_bs.append(["Reserves and Surplus", "2", 5000000.00, 4500000.00])
    ws_bs.append(["TOTAL LIABILITIES & EQUITY", "", 201403570.73, 180000000.00])
    ws_bs.append(["II. ASSETS", "", "", ""])
    ws_bs.append(["Fixed Assets", "3", 80000000.00, 75000000.00])
    ws_bs.append(["Current Assets", "4", 121402126.89, 105000000.00])
    ws_bs.append(["TOTAL ASSETS", "", 201402126.89, 180000000.00])
    
    # Profit & Loss
    ws_pl = wb.create_sheet(title="Profit & Loss")
    ws_pl.append(["Zion Lifts LLP - Statement of Profit and Loss for the year ended 31st March 2026"])
    ws_pl.append(["Particulars", "Note No.", "Figures for 2025-26", "Figures for 2024-25"])
    ws_pl.append(["Revenue from operations", "19", 50000000.00, 45000000.00])
    ws_pl.append(["Total Revenue", "", 50000000.00, 45000000.00])
    
    # Notes with formula range issue
    ws_notes = wb.create_sheet(title="Notes to BS")
    ws_notes.append(["Notes forming part of financial statements"])
    for i in range(2, 30):
        ws_notes.append([f"Item {i}", 1000.0 * i])
    # Formula that omits rows
    ws_notes["D51"] = "=SUM(D30:D49)"
    
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# 1. User can send a message and assistant returns non-empty response
def test_user_can_send_message_and_get_non_empty_response():
    headers = get_auth_headers()
    conv_id = "test_conv_basic_flow"
    
    res = client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={
            "conversation_id": conv_id,
            "message": "What is the role of debit and credit in accounting?",
        },
    )
    assert res.status_code == 200
    data = res.json()
    assert data["message"] is not None
    assert len(data["message"].strip()) > 0
    assert data["conversation_id"] == conv_id


# 2 & 3. Uploaded document context is loaded
def test_upload_and_document_context_loaded():
    headers = get_auth_headers()
    wb_bytes = create_sample_financial_workbook_bytes()
    conv_id = "test_conv_doc_loaded"

    upload_res = client.post(
        "/api/v1/assistant/upload",
        headers=headers,
        files={"file": ("Zion_Sample_Fin.xlsx", wb_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"conversation_id": conv_id},
    )
    assert upload_res.status_code == 200
    doc_id = upload_res.json()["document_id"]

    # Send document question
    chat_res = client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={
            "conversation_id": conv_id,
            "message": "What is this file?",
            "file_ids": [doc_id],
        },
    )
    assert chat_res.status_code == 200
    chat_data = chat_res.json()
    assert "Zion_Sample_Fin.xlsx" in chat_data["message"] or "Balance Sheet" in chat_data["message"]
    assert chat_data["intent"] == "DOCUMENT_METADATA"


# 4 & 5. Conversation & Assistant response persist in database
def test_conversation_and_messages_persist_in_db():
    headers = get_auth_headers()
    conv_id = "test_conv_persisted_check"

    res = client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={
            "conversation_id": conv_id,
            "message": "Explain accrued expenses.",
        },
    )
    assert res.status_code == 200

    db = SessionLocal()
    try:
        conv = db.query(AssistantConversation).filter(AssistantConversation.id == conv_id).first()
        assert conv is not None
        
        # Check messages
        msgs = db.query(AssistantMessage).filter(AssistantMessage.conversation_id == conv_id).all()
        assert len(msgs) >= 2  # user message and assistant message
        user_msg = next((m for m in msgs if m.role == "user"), None)
        asst_msg = next((m for m in msgs if m.role == "assistant"), None)
        assert user_msg is not None and "accrued expenses" in user_msg.content.lower()
        assert asst_msg is not None and len(asst_msg.content) > 0
    finally:
        db.close()


# 6. Multiple questions maintain context
def test_multiple_questions_maintain_context():
    headers = get_auth_headers()
    wb_bytes = create_sample_financial_workbook_bytes()
    conv_id = "test_conv_multi_turn_persisted"

    upload_res = client.post(
        "/api/v1/assistant/upload",
        headers=headers,
        files={"file": ("Zion_Sample_Fin.xlsx", wb_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"conversation_id": conv_id},
    )
    doc_id = upload_res.json()["document_id"]

    # Turn 1: What is wrong?
    r1 = client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={
            "conversation_id": conv_id,
            "message": "What is wrong with this balance sheet?",
            "file_ids": [doc_id],
        },
    )
    assert r1.status_code == 200
    assert "1,443.84" in r1.json()["message"]

    # Turn 2: What is the difference? (Contextual follow-up without resending file_ids)
    r2 = client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={
            "conversation_id": conv_id,
            "message": "What is the difference?",
        },
    )
    assert r2.status_code == 200
    assert "1,443.84" in r2.json()["message"]


# 7. Finding question returns relevant deterministic finding
def test_finding_question_returns_deterministic_finding():
    headers = get_auth_headers()
    wb_bytes = create_sample_financial_workbook_bytes()
    conv_id = "test_conv_findings_check"

    upload_res = client.post(
        "/api/v1/assistant/upload",
        headers=headers,
        files={"file": ("Zion_Sample_Fin.xlsx", wb_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"conversation_id": conv_id},
    )
    doc_id = upload_res.json()["document_id"]

    res = client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={
            "conversation_id": conv_id,
            "message": "Where is the mismatch?",
            "file_ids": [doc_id],
        },
    )
    assert res.status_code == 200
    data = res.json()
    assert len(data["findings"]) > 0
    assert "1,443.84" in data["message"] or any("1,443.84" in str(f.get("difference", "")) for f in data["findings"])


# 8. Formula question returns formula findings
def test_formula_question_returns_formula_findings():
    headers = get_auth_headers()
    wb_bytes = create_sample_financial_workbook_bytes()
    conv_id = "test_conv_formula_check"

    upload_res = client.post(
        "/api/v1/assistant/upload",
        headers=headers,
        files={"file": ("Zion_Sample_Fin.xlsx", wb_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"conversation_id": conv_id},
    )
    doc_id = upload_res.json()["document_id"]

    res = client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={
            "conversation_id": conv_id,
            "message": "Are there formula errors or range errors in this workbook?",
            "file_ids": [doc_id],
        },
    )
    assert res.status_code == 200
    data = res.json()
    assert data["intent"] == "FORMULA_QUERY"
    assert "formula" in data["message"].lower() or "range" in data["message"].lower()


# 9. Reconciliation question returns reconciliation data
def test_reconciliation_question():
    headers = get_auth_headers()
    conv_id = "test_conv_recon"

    res = client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={
            "conversation_id": conv_id,
            "message": "How many transactions are unmatched in reconciliation?",
        },
    )
    assert res.status_code == 200
    data = res.json()
    assert data["intent"] == "RECONCILIATION_QUERY"
    assert "reconciliation" in data["message"].lower()


# 10. Unknown question returns safe response
def test_unknown_question_safe_response():
    headers = get_auth_headers()
    conv_id = "test_conv_unknown"

    res = client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={
            "conversation_id": conv_id,
            "message": "Xyzzwq abracadabra random tokens 123456",
        },
    )
    assert res.status_code == 200
    data = res.json()
    assert data["message"] is not None
    assert len(data["message"].strip()) > 0
    assert not data["message"].startswith("{")  # Never raw JSON


# 11 & 12. AI provider fallback and timeout safety
def test_ai_provider_fallback_on_exception(monkeypatch):
    headers = get_auth_headers()
    wb_bytes = create_sample_financial_workbook_bytes()
    conv_id = "test_conv_fallback"

    upload_res = client.post(
        "/api/v1/assistant/upload",
        headers=headers,
        files={"file": ("Zion_Sample_Fin.xlsx", wb_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"conversation_id": conv_id},
    )
    doc_id = upload_res.json()["document_id"]

    async def faulty_generate(self, prompt, system_prompt=None, context=None):
        raise RuntimeError("OpenAI API unreachable or timeout")

    monkeypatch.setattr(MockAIProvider, "generate_response", faulty_generate)

    res = client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={
            "conversation_id": conv_id,
            "message": "Explain this balance sheet simply.",
            "file_ids": [doc_id],
        },
    )
    assert res.status_code == 200
    data = res.json()
    assert "couldn't generate the ai explanation" in data["message"].lower() or "unavailable" in data["message"].lower()
    # Must still contain deterministic facts
    assert "1,443.84" in data["message"]


# 13 & 14. Unauthorized conversation & document return 403
def test_unauthorized_conversation_access_returns_403():
    old_override = fastapi_app.dependency_overrides.pop(get_current_user, None)
    try:
        import uuid
        user1_headers = get_auth_headers("user_alpha@litebooks.ai")
        user2_headers = get_other_user_headers("user_beta@litebooks.ai")
        conv_id = f"alpha_private_conv_{uuid.uuid4().hex[:8]}"

        # User 1 creates conversation
        r1 = client.post(
            "/api/v1/assistant/chat",
            headers=user1_headers,
            json={
                "conversation_id": conv_id,
                "message": "Hello from Alpha",
            },
        )
        assert r1.status_code == 200

        # User 2 tries to access User 1's conversation
        r2 = client.post(
            "/api/v1/assistant/chat",
            headers=user2_headers,
            json={
                "conversation_id": conv_id,
                "message": "Intruder message",
            },
        )
        assert r2.status_code == 403

        # User 2 tries to GET conversation context
        r3 = client.get(
            f"/api/v1/assistant/conversation/{conv_id}",
            headers=user2_headers,
        )
        assert r3.status_code == 403
    finally:
        if old_override:
            fastapi_app.dependency_overrides[get_current_user] = old_override


# 15. Empty document context is handled gracefully
def test_empty_document_context_handled():
    headers = get_auth_headers()
    conv_id = "test_conv_no_doc"

    res = client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={
            "conversation_id": conv_id,
            "message": "What is wrong with this workbook?",
            # No files attached
        },
    )
    assert res.status_code == 200
    data = res.json()
    assert "upload" in data["message"].lower() or "attach" in data["message"].lower() or "no document" in data["message"].lower() or "financial" in data["message"].lower()


# 16. MockAIProvider never returns blank response
def test_mock_ai_provider_never_blank():
    import asyncio
    provider = MockAIProvider()
    resp = asyncio.run(provider.generate_response(
        prompt="What is wrong with the balance sheet?",
        context={
            "document": {"name": "Test.xlsx"},
            "findings": [{"type": "AMOUNT_MISMATCH", "severity": "HIGH", "expected": 200, "actual": 250, "difference": 50}],
        }
    ))
    assert resp is not None
    assert len(resp.strip()) > 0
    assert "50" in resp or "₹" in resp


# 17. Correction requests never auto-approve
def test_correction_never_auto_approves():
    headers = get_auth_headers()
    wb_bytes = create_sample_financial_workbook_bytes()
    conv_id = "test_conv_no_auto_approve"

    upload_res = client.post(
        "/api/v1/assistant/upload",
        headers=headers,
        files={"file": ("Zion_Sample_Fin.xlsx", wb_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"conversation_id": conv_id},
    )
    doc_id = upload_res.json()["document_id"]

    res = client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={
            "conversation_id": conv_id,
            "message": "Fix it immediately.",
            "file_ids": [doc_id],
        },
    )
    assert res.status_code == 200
    data = res.json()
    assert data["intent"] == "CORRECTION_REQUEST"
    # Must advise review/approval, not execute
    assert "approve" in data["message"].lower() or "review" in data["message"].lower() or "candidate" in data["message"].lower()


# 18. Financial values are grounded in deterministic context
def test_values_grounded_in_deterministic_context():
    headers = get_auth_headers()
    wb_bytes = create_sample_financial_workbook_bytes()
    conv_id = "test_conv_grounding"

    upload_res = client.post(
        "/api/v1/assistant/upload",
        headers=headers,
        files={"file": ("Zion_Sample_Fin.xlsx", wb_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"conversation_id": conv_id},
    )
    doc_id = upload_res.json()["document_id"]

    res = client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={
            "conversation_id": conv_id,
            "message": "What is wrong with this balance sheet?",
            "file_ids": [doc_id],
        },
    )
    assert res.status_code == 200
    data = res.json()
    # Expect deterministic difference 1443.84
    assert "1,443.84" in data["message"]


# 19. Original file remains immutable
def test_original_file_remains_immutable():
    headers = get_auth_headers()
    wb_bytes = create_sample_financial_workbook_bytes()
    orig_sha256 = hashlib.sha256(wb_bytes).hexdigest()
    conv_id = "test_conv_immutable"

    upload_res = client.post(
        "/api/v1/assistant/upload",
        headers=headers,
        files={"file": ("Zion_Sample_Fin.xlsx", wb_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"conversation_id": conv_id},
    )
    doc_id = upload_res.json()["document_id"]

    # Multiple chats including correction attempt
    client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={"conversation_id": conv_id, "message": "What is wrong?", "file_ids": [doc_id]},
    )
    client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={"conversation_id": conv_id, "message": "Fix this mismatch."},
    )

    # Verify original file in DB still matches checksum
    db = SessionLocal()
    try:
        doc = db.query(Document).filter(Document.id == doc_id).first()
        assert doc is not None
        assert doc.checksum == orig_sha256 or doc.original_sha256 == orig_sha256
    finally:
        db.close()
