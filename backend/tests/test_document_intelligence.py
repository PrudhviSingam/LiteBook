import io
import math
from pathlib import Path
import openpyxl
import pytest

from app.modules.reconciliation.financial_validator import FinancialValidator, format_currency_inr
from app.modules.reconciliation.matching import match_transactions
from app.modules.reconciliation.schemas import (
    DiscrepancySeverityEnum,
    DiscrepancyTypeEnum,
    DocumentTypeEnum,
    NormalizedTransaction,
    ReconciliationRuleConfig,
)
from app.modules.reconciliation.workbook_inspector import (
    DocumentClassifier,
    FinancialPeriodDetector,
    WorkbookInspector,
    format_indian_currency,
    parse_indian_currency,
    parse_indian_number,
)


# =============================================================================
# 1. DOCUMENT CLASSIFICATION (11 TYPES + DETERMINISTIC EVIDENCE)
# =============================================================================

def test_document_classification_balance_sheet():
    corpus = "Statement of Assets and Liabilities Particulars Partner Capital Accounts Reserves and Surplus Current Liabilities Trade Payables Total Liabilities Tangible Fixed Assets Inventories Trade Receivables Cash and Bank Balances Total Assets"
    doc_type, evidence = DocumentClassifier.classify(corpus, "Balance Sheet")
    assert doc_type == DocumentTypeEnum.BALANCE_SHEET
    assert evidence.confidence == "HIGH"
    assert len(evidence.signals_found) >= 2
    assert "Balance Sheet" in evidence.explanation or "balance sheet" in evidence.explanation.lower()


def test_document_classification_profit_loss():
    corpus = "Statement of Profit and Loss Revenue from Operations Other Income Total Expenses Net Profit Before Tax EBITDA"
    doc_type, evidence = DocumentClassifier.classify(corpus, "Profit & Loss")
    assert doc_type == DocumentTypeEnum.PROFIT_LOSS
    assert evidence.confidence == "HIGH"
    assert any("profit and loss" in s or "revenue" in s for s in evidence.signals_found)


def test_document_classification_trial_balance():
    corpus = "Trial Balance as on 31-03-2026 Ledger Account Opening Balance Debit Credit Closing Balance Total"
    doc_type, evidence = DocumentClassifier.classify(corpus, "Trial Balance")
    assert doc_type == DocumentTypeEnum.TRIAL_BALANCE
    assert evidence.confidence in ("HIGH", "MEDIUM")


def test_document_classification_cash_flow():
    corpus = "Cash Flow Statement Cash flow from Operating Activities Investing Activities Financing Activities Net Increase in Cash"
    doc_type, evidence = DocumentClassifier.classify(corpus, "Cash Flow")
    assert doc_type == DocumentTypeEnum.CASH_FLOW
    assert evidence.confidence == "HIGH"


def test_document_classification_bank_statement():
    corpus = "HDFC Bank Current Account Statement Date Value Date Chq No Description Withdrawal Deposit Balance"
    doc_type, evidence = DocumentClassifier.classify(corpus, "Bank Statement", has_withdrawal_deposit=True)
    assert doc_type == DocumentTypeEnum.BANK_STATEMENT
    assert evidence.confidence == "HIGH"


def test_document_classification_accounting_ledger():
    corpus = "Account Ledger Date Particulars Vch No Debit Credit Balance Narration"
    doc_type, evidence = DocumentClassifier.classify(corpus, "Vendor Ledger")
    assert doc_type == DocumentTypeEnum.ACCOUNTING_LEDGER
    assert evidence.confidence == "HIGH"


def test_document_classification_general_ledger():
    corpus = "General Ledger Chart of Accounts GL Code Account Name Opening Balance Debit Credit Net Balance"
    doc_type, evidence = DocumentClassifier.classify(corpus, "General Ledger")
    assert doc_type == DocumentTypeEnum.GENERAL_LEDGER
    assert evidence.confidence == "HIGH"


def test_document_classification_tax_document():
    corpus = "GSTR-3B Return Summary GSTIN 27AABCT1234F1Z5 Taxable Value CGST SGST IGST Tax Invoice TDS Certificate Form 26AS"
    doc_type, evidence = DocumentClassifier.classify(corpus, "GST Return")
    assert doc_type == DocumentTypeEnum.TAX_DOCUMENT
    assert evidence.confidence == "HIGH"


def test_document_classification_unknown():
    corpus = "Employee cafeteria menu Monday sandwich Tuesday soup Wednesday salad"
    doc_type, evidence = DocumentClassifier.classify(corpus, "Menu")
    assert doc_type == DocumentTypeEnum.UNKNOWN
    assert evidence.confidence == "LOW"


# =============================================================================
# 2. FINANCIAL PERIOD DETECTION & NORMALIZATION
# =============================================================================

def test_period_detection_indian_fy():
    p1 = FinancialPeriodDetector.detect_period(["Financial Statements for FY 2025-26"])
    assert p1.period_type == "FINANCIAL_YEAR"
    assert p1.financial_year == "2025-26"
    assert p1.start_date == "2025-04-01"
    assert p1.end_date == "2026-03-31"

    p2 = FinancialPeriodDetector.detect_period(["Period 01/04/2025 - 31/03/2026"])
    assert p2.financial_year == "2025-26"
    assert p2.start_date == "2025-04-01"
    assert p2.end_date == "2026-03-31"

    p3 = FinancialPeriodDetector.detect_period(["For the year ended 31 March 2026"])
    assert p3.financial_year == "2025-26"
    assert p3.end_date == "2026-03-31"

    p4 = FinancialPeriodDetector.detect_period(["Apr 2025 - Mar 2026"])
    assert p4.financial_year == "2025-26"


def test_period_detection_from_filename():
    p = FinancialPeriodDetector.detect_period([], filename="Zion_Lifts_Audit_FY2025-26.xlsx")
    assert p.financial_year == "2025-26"
    assert p.start_date == "2025-04-01"
    assert p.end_date == "2026-03-31"


# =============================================================================
# 3. INDIAN NUMBER PARSING & UNITS
# =============================================================================

def test_indian_number_parsing_variations():
    # Indian comma notation with Rupee symbol
    assert parse_indian_currency("₹20,14,02,126.89") == 201402126.89
    # Without rupee symbol
    assert parse_indian_currency("20,14,02,126.89") == 201402126.89
    # Parentheses negative
    assert parse_indian_currency("(1,443.84)") == -1443.84
    # Hyphen negative
    assert parse_indian_currency("-1,250.00") == -1250.00
    # Rs. prefix
    assert parse_indian_currency("Rs. 1,250") == 1250.0
    assert parse_indian_currency("INR 5,00,000") == 500000.0
    # Percentages
    assert parse_indian_currency("18%") == 0.18
    assert parse_indian_currency("18.00%") == 0.18
    # Formatting helper
    assert format_indian_currency(201402126.89) == "₹20,14,02,126.89"
    assert format_indian_currency(-1443.84) == "-₹1,443.84"


def test_parse_indian_number_provenance():
    val, raw, unit = parse_indian_number("₹20,14,02,126.89")
    assert val == 201402126.89
    assert raw == "₹20,14,02,126.89"
    assert unit == "INR"


# =============================================================================
# 4. DETERMINISTIC FINANCIAL VALIDATION RULES
# =============================================================================

def test_imbalanced_balance_sheet_deterministic_calculation():
    """Validates the exact real-world scenario specified in Phase 19:

    Total Assets: ₹20,14,02,126.89
    Total Liabilities: ₹20,14,03,570.73
    Expected difference: exactly 1,443.84 computed deterministically.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Balance Sheet"

    ws.append(["STATEMENT OF ASSETS AND LIABILITIES AS AT 31 MARCH 2026", None, None])
    ws.append([])
    ws.append(["Particulars", "Note", "31 March 2026"])
    ws.append(["I. EQUITY AND LIABILITIES", None, None])
    ws.append(["Partner Capital", "1", 50000000.0])
    ws.append(["Trade Payables", "2", 151403570.73])
    ws.append(["TOTAL LIABILITIES", None, 201403570.73])
    ws.append([])
    ws.append(["II. ASSETS", None, None])
    ws.append(["Fixed Assets", "3", 100000000.0])
    ws.append(["Current Assets", "4", 101402126.89])
    ws.append(["TOTAL ASSETS", None, 201402126.89])

    discrepancies = FinancialValidator.validate_balance_sheets(wb)
    assert len(discrepancies) == 1
    d = discrepancies[0]
    assert d.type == DiscrepancyTypeEnum.BALANCE_SHEET_IMBALANCE
    assert d.severity == DiscrepancySeverityEnum.CRITICAL
    # Verifies difference is exactly 1443.84
    assert math.isclose(d.difference, 1443.84, abs_tol=0.001)
    assert "Liabilities and equity exceed total assets by ₹1,443.84" in d.description
    assert d.cell_address == "C12"
    assert "Possible cause" in d.possible_causes[0]


def test_balanced_balance_sheet_no_discrepancy():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Balance Sheet"
    ws.append(["Particulars", "Note", "31 March 2026"])
    ws.append(["TOTAL LIABILITIES", None, 5000000.0])
    ws.append(["TOTAL ASSETS", None, 5000000.0])

    discrepancies = FinancialValidator.validate_balance_sheets(wb)
    assert len(discrepancies) == 0


def test_imbalanced_trial_balance_validation():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trial Balance"
    ws.append(["Particulars", "Debit", "Credit"])
    ws.append(["Cash", 100000.0, None])
    ws.append(["Capital", None, 95000.0])
    ws.append(["Grand Total", 100000.0, 95000.0])

    discrepancies = FinancialValidator.validate_trial_balance(wb)
    assert len(discrepancies) == 1
    d = discrepancies[0]
    assert d.type == DiscrepancyTypeEnum.TRIAL_BALANCE_IMBALANCE
    assert math.isclose(d.difference, 5000.0, abs_tol=0.01)


def test_profit_loss_calculation_validation():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Profit & Loss"
    ws.append(["Particulars", "FY 2025-26"])
    ws.append(["Revenue from Operations", 1000000.0])
    ws.append(["Total Expenses", 800000.0])
    # Expected profit is 200,000, but workbook says 250,000 (mismatch)
    ws.append(["Net Profit", 250000.0])

    discrepancies = FinancialValidator.validate_profit_loss(wb)
    assert len(discrepancies) == 1
    d = discrepancies[0]
    assert d.type == DiscrepancyTypeEnum.AMOUNT_MISMATCH
    assert math.isclose(d.difference, 50000.0, abs_tol=0.01)


# =============================================================================
# 5. FORMULA AWARENESS
# =============================================================================

def test_formula_awareness_range_omission():
    wb_data = openpyxl.Workbook()
    ws_d = wb_data.active
    ws_d.title = "Balance Sheet"
    ws_d.append(["Particulars", "Amount"])
    ws_d.append(["Asset A", 100.0])
    ws_d.append(["Asset B", 200.0])
    ws_d.append(["Asset C", 300.0])
    ws_d.append(["TOTAL ASSETS", 600.0])

    # Formula workbook where formula only sums up to row 3 instead of 4
    wb_form = openpyxl.Workbook()
    ws_f = wb_form.active
    ws_f.title = "Balance Sheet"
    ws_f.append(["Particulars", "Amount"])
    ws_f.append(["Asset A", 100.0])
    ws_f.append(["Asset B", 200.0])
    ws_f.append(["Asset C", 300.0])
    ws_f.append(["TOTAL ASSETS", "=SUM(B2:B3)"])  # Omits row 4

    discrepancies = FinancialValidator.validate_formulas(wb_data, wb_form)
    assert any(d.type == DiscrepancyTypeEnum.FORMULA_ERROR and "omits" in d.description for d in discrepancies)


# =============================================================================
# 6. RECONCILIATION MATCHING TIERS & CONFIDENCE
# =============================================================================

def test_reconciliation_deterministic_tiers():
    b1 = NormalizedTransaction(
        id="b1", date="2026-03-15", description="Apex Payment", reference="REF101",
        debit=5000.0, credit=0.0, amount=5000.0, direction="debit", source="bank"
    )
    l1 = NormalizedTransaction(
        id="l1", date="2026-03-15", description="Apex Payment", reference="REF101",
        debit=5000.0, credit=0.0, amount=5000.0, direction="debit", source="ledger"
    )

    # Tier 1: Exact reference + exact amount + date
    results, summary = match_transactions([b1], [l1], ReconciliationRuleConfig(date_tolerance_days=1))
    assert len(results) == 1
    assert results[0].status == "MATCHED"
    assert results[0].confidence == "high"
    assert results[0].match_tier == 1

    # Tier 2: Exact amount + date within ±1 day (no matching reference)
    b2 = NormalizedTransaction(
        id="b2", date="2026-03-15", description="Electricity Bill", reference="",
        debit=3200.0, credit=0.0, amount=3200.0, direction="debit", source="bank"
    )
    l2 = NormalizedTransaction(
        id="l2", date="2026-03-16", description="Electricity Bill", reference="",
        debit=3200.0, credit=0.0, amount=3200.0, direction="debit", source="ledger"
    )
    results2, _ = match_transactions([b2], [l2], ReconciliationRuleConfig(date_tolerance_days=1))
    assert len(results2) == 1
    assert results2[0].status == "MATCHED"
    assert results2[0].match_tier in (2, 3)

    # Tier 4: Date mismatch beyond tolerance
    b3 = NormalizedTransaction(
        id="b3", date="2026-03-01", description="Vendor Advance", reference="ADV99",
        debit=15000.0, credit=0.0, amount=15000.0, direction="debit", source="bank"
    )
    l3 = NormalizedTransaction(
        id="l3", date="2026-03-25", description="Vendor Advance", reference="ADV99",
        debit=15000.0, credit=0.0, amount=15000.0, direction="debit", source="ledger"
    )
    results3, _ = match_transactions([b3], [l3], ReconciliationRuleConfig(date_tolerance_days=1))
    assert len(results3) == 1
    assert results3[0].status == "DATE_MISMATCH"
    assert results3[0].match_tier == 4


# =============================================================================
# 7. REAL-WORLD WORKBOOK TEST (ZION LIFTS FINANCIALS)
# =============================================================================

def test_real_zion_lifts_workbook_analysis():
    """Inspects the actual Zion_Lifts_Financials.xlsx file from the workspace.

    Verifies statement structure, sheets, period, and full provenance.
    """
    wb_path = Path("Zion_Lifts_Financials.xlsx")
    assert wb_path.exists(), "Zion_Lifts_Financials.xlsx must exist in backend directory"

    inspection = WorkbookInspector.inspect(wb_path, "Zion_Lifts_Financials.xlsx")
    assert inspection.total_sheets == 4
    assert inspection.document_type == DocumentTypeEnum.FINANCIAL_STATEMENT
    assert inspection.period_info is not None
    assert inspection.period_info.financial_year == "2025-26"
    assert inspection.period_info.start_date == "2025-04-01"
    assert inspection.period_info.end_date == "2026-03-31"

    # Check that structured tables exist with exact cell coordinates
    assert len(inspection.tables) >= 1
    bs_table = next((t for t in inspection.tables if t.sheet_name == "Balance Sheet"), None)
    assert bs_table is not None
    assert bs_table.table_type == DocumentTypeEnum.BALANCE_SHEET
    assert len(bs_table.rows) > 0

    # Verify provenance for a row (e.g. Partner Capital Accounts)
    cap_row = next((r for r in bs_table.rows if "Capital" in r.label), None)
    assert cap_row is not None
    assert "31 March 2026" in cap_row.source_cells
    cell_prov = cap_row.source_cells["31 March 2026"]
    assert cell_prov.sheet_name == "Balance Sheet"
    assert cell_prov.cell_address.startswith("C")
    assert cell_prov.normalized_value == 5000000.0
