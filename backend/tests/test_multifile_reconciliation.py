import asyncio
from pathlib import Path
import uuid
import pytest
from sqlalchemy.orm import Session

from app.core.database import SessionLocal, engine, Base
import app.models  # Register all models
Base.metadata.create_all(bind=engine)

from app.models.document import Document
from app.modules.reconciliation.schemas import (
    DocumentTypeEnum,
    ReconciliationStatusEnum,
    ReconciliationRuleConfig,
)
from app.modules.reconciliation.normalization import normalize_transactions
from app.modules.reconciliation.services import ReconciliationService
from app.modules.ai_assistant.assistant_service import AssistantService
from app.modules.ai_assistant.schemas import ChatRequest, AssistantIntentEnum


@pytest.fixture
def db():
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()


@pytest.fixture
def test_files(db: Session):
    """Provides or registers LiteBooks_Reconciliation_Test_Data.xlsx and bankFile.xlsx."""
    import openpyxl
    storage_dir = Path("storage/uploads")
    storage_dir.mkdir(parents=True, exist_ok=True)

    ledger_path = storage_dir / "LiteBooks_Reconciliation_Test_Data.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "General Ledger"
    ws.append(["Date", "Description", "Reference", "Amount"])
    ws.append(["2026-09-01", "Vendor Supplies", "TXN1001", 10000.0])
    ws.append(["2026-09-02", "Software Subscription", "TXN1002", 4500.0])
    ws.append(["2026-09-03", "Office Rent", "TXN1003", 25000.0])
    ws.append(["2026-09-04", "Consulting Revenue", "TXN1004", 12000.0])
    ws.append(["2026-09-05", "Utility Electric", "TXN1005", 1800.0])
    ws.append(["2026-09-06", "Hardware Repair", "TXN1006", 3200.0])
    wb.save(ledger_path)

    ledger_doc = db.query(Document).filter(
        Document.original_filename == "LiteBooks_Reconciliation_Test_Data.xlsx"
    ).order_by(Document.id.desc()).first()

    if not ledger_doc:
        ledger_doc = Document(
            stored_filename=f"test_data_{uuid.uuid4().hex[:8]}.xlsx",
            original_filename="LiteBooks_Reconciliation_Test_Data.xlsx",
            file_path=str(ledger_path),
            file_size=ledger_path.stat().st_size,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            category="ledger",
        )
        db.add(ledger_doc)
        db.commit()
        db.refresh(ledger_doc)
    else:
        ledger_doc.file_path = str(ledger_path)
        db.commit()

    bank_path = storage_dir / "bankFile.xlsx"
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "Sheet1"
    ws2.append(["Date", "Description", "Reference", "Withdrawal", "Deposit"])
    ws2.append(["2026-09-01", "Vendor Supplies", "TXN1001", 10000.0, 0.0])
    ws2.append(["2026-09-02", "Software Subscription", "TXN1002", 5400.0, 0.0])
    ws2.append(["2026-09-03", "Office Rent", "TXN1003", 25000.0, 0.0])
    ws2.append(["2026-09-05", "Utility Electric", "TXN1005", 1800.0, 0.0])
    ws2.append(["2026-09-07", "Bank Service Charges", "TXN1007", 450.0, 0.0])
    wb2.save(bank_path)

    bank_doc = db.query(Document).filter(
        Document.original_filename == "bankFile.xlsx"
    ).order_by(Document.id.desc()).first()

    if not bank_doc:
        bank_doc = Document(
            stored_filename=f"test_bank_{uuid.uuid4().hex[:8]}.xlsx",
            original_filename="bankFile.xlsx",
            file_path=str(bank_path),
            file_size=bank_path.stat().st_size,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            category="bank_statement",
        )
        db.add(bank_doc)
        db.commit()
        db.refresh(bank_doc)
    else:
        bank_doc.file_path = str(bank_path)
        db.commit()

    return {"ledger": ledger_doc, "bank": bank_doc}


# =============================================================================
# 1. Deterministic Source Discovery Tests
# =============================================================================

def test_source_discovery_both_files(db: Session, test_files):
    ledger_doc = test_files["ledger"]
    bank_doc = test_files["bank"]

    res = ReconciliationService.discover_sources(db, [ledger_doc.id, bank_doc.id])
    assert res["status"] == "ready"
    assert res["selected_ledger"] is not None
    assert res["selected_bank"] is not None
    assert res["selected_ledger"]["document_id"] == ledger_doc.id
    assert res["selected_ledger"]["sheet_name"] == "General Ledger"
    assert res["selected_bank"]["document_id"] == bank_doc.id
    assert res["selected_bank"]["sheet_name"] == "Sheet1"


def test_source_discovery_missing_bank(db: Session, test_files):
    ledger_doc = test_files["ledger"]
    res = ReconciliationService.discover_sources(db, [ledger_doc.id])
    assert res["status"] == "REVIEW_REQUIRED"
    assert res["reason"] == "BANK_STATEMENT_NOT_FOUND"
    assert len(res["ledger_candidates"]) >= 1
    assert len(res["bank_candidates"]) == 0


def test_source_discovery_missing_ledger(db: Session, test_files):
    bank_doc = test_files["bank"]
    res = ReconciliationService.discover_sources(db, [bank_doc.id])
    assert res["status"] == "REVIEW_REQUIRED"
    assert res["reason"] == "GENERAL_LEDGER_NOT_FOUND"
    assert len(res["ledger_candidates"]) == 0
    assert len(res["bank_candidates"]) >= 1


# =============================================================================
# 2. Normalization & Complete Provenance Tests
# =============================================================================

def test_normalization_and_provenance(test_files):
    ledger_doc = test_files["ledger"]
    bank_doc = test_files["bank"]

    ledger_txs = normalize_transactions(
        Path(ledger_doc.file_path),
        filename=ledger_doc.original_filename,
        sheet_name="General Ledger",
        source="ledger",
        document_id=ledger_doc.id,
        source_type="GENERAL_LEDGER",
    )

    bank_txs = normalize_transactions(
        Path(bank_doc.file_path),
        filename=bank_doc.original_filename,
        sheet_name="Sheet1",
        source="bank",
        document_id=bank_doc.id,
        source_type="BANK_STATEMENT",
    )

    assert len(ledger_txs) == 6
    assert len(bank_txs) == 5

    # Check complete provenance on ledger items
    for tx in ledger_txs:
        assert tx.document_id == ledger_doc.id
        assert tx.filename == "LiteBooks_Reconciliation_Test_Data.xlsx"
        assert tx.sheet_name == "General Ledger"
        assert tx.source_type == "GENERAL_LEDGER"
        assert tx.original_row_number is not None
        assert tx.original_row_number >= 2

    # Check complete provenance on bank items
    for tx in bank_txs:
        assert tx.document_id == bank_doc.id
        assert tx.filename == "bankFile.xlsx"
        assert tx.sheet_name == "Sheet1"
        assert tx.source_type == "BANK_STATEMENT"
        assert tx.original_row_number is not None
        assert tx.original_row_number >= 2


# =============================================================================
# 3. Deterministic Multi-File Reconciliation Matching Tests
# =============================================================================

def test_deterministic_reconciliation_math(db: Session, test_files):
    ledger_doc = test_files["ledger"]
    bank_doc = test_files["bank"]

    recon_res = asyncio.run(ReconciliationService.run_reconciliation(
        db=db,
        bank_document_id=bank_doc.id,
        ledger_document_id=ledger_doc.id,
        bank_sheet_name="Sheet1",
        ledger_sheet_name="General Ledger",
        rules=ReconciliationRuleConfig(date_tolerance_days=2, amount_tolerance=0.0),
    ))

    summary = recon_res.summary
    assert summary.total_ledger_transactions == 6
    assert summary.total_bank_transactions == 5
    assert summary.matched_count == 3
    assert summary.amount_mismatch_count == 1
    assert summary.missing_in_bank_count == 2
    assert summary.missing_in_ledger_count == 1
    assert summary.date_mismatch_count == 0
    assert summary.total_variance == 13850.0
    assert summary.reconciliation_rate == 50.0

    # Verify amount mismatch specifically
    amt_item = next(t for t in recon_res.transactions if t.status == ReconciliationStatusEnum.AMOUNT_MISMATCH)
    assert amt_item.reference == "TXN1002"
    assert amt_item.ledger_amount == 4500.0
    assert amt_item.bank_amount == 5400.0
    assert amt_item.difference == 900.0

    # Verify missing in bank
    missing_bank_items = [t for t in recon_res.transactions if t.status == ReconciliationStatusEnum.MISSING_IN_BANK]
    assert len(missing_bank_items) == 2
    mb_refs = {t.reference for t in missing_bank_items}
    assert mb_refs == {"TXN1004", "TXN1006"}

    # Verify missing in ledger
    missing_ledger_items = [t for t in recon_res.transactions if t.status == ReconciliationStatusEnum.MISSING_IN_LEDGER]
    assert len(missing_ledger_items) == 1
    assert missing_ledger_items[0].reference == "TXN1007"
    assert missing_ledger_items[0].bank_amount == 450.0


# =============================================================================
# 4. Discrepancy Findings Formatting Tests
# =============================================================================

def test_format_findings_for_assistant(db: Session, test_files):
    ledger_doc = test_files["ledger"]
    bank_doc = test_files["bank"]

    recon_res = asyncio.run(ReconciliationService.run_reconciliation(
        db=db,
        bank_document_id=bank_doc.id,
        ledger_document_id=ledger_doc.id,
        bank_sheet_name="Sheet1",
        ledger_sheet_name="General Ledger",
    ))

    findings = ReconciliationService.format_findings_for_assistant(
        items=recon_res.transactions,
        ledger_source={"document_id": ledger_doc.id, "filename": ledger_doc.original_filename, "sheet_name": "General Ledger"},
        bank_source={"document_id": bank_doc.id, "filename": bank_doc.original_filename, "sheet_name": "Sheet1"},
    )

    # 1 amount mismatch + 2 missing in bank + 1 missing in ledger = 4 discrepancy findings
    assert len(findings) == 4

    f_amt = next(f for f in findings if f.finding_type == "AMOUNT_MISMATCH")
    assert f_amt.ledger_evidence is not None
    assert f_amt.bank_evidence is not None
    assert f_amt.ledger_evidence["row_number"] == 3
    assert f_amt.bank_evidence["row_number"] == 3
    assert f_amt.difference == "₹900.00"
    assert f_amt.confidence == "HIGH"


# =============================================================================
# 5. Assistant Chat Data Pipeline & Grounded Accounting Answers
# =============================================================================

def test_assistant_reconciliation_amount_mismatches(db: Session, test_files):
    ledger_doc = test_files["ledger"]
    bank_doc = test_files["bank"]

    req = ChatRequest(
        message="Find any amount mismatches between the General Ledger and Bank Statement.",
        file_ids=[ledger_doc.id, bank_doc.id],
    )
    resp = asyncio.run(AssistantService.process_chat(db, req))

    assert resp.confidence == "HIGH"
    assert len(resp.findings) > 0
    msg = resp.message
    assert "TXN1002" in msg
    assert "4,500" in msg
    assert "5,400" in msg
    assert "900" in msg
    assert "transposition" in msg.lower() or "data entry" in msg.lower()


def test_assistant_reconciliation_general_difference(db: Session, test_files):
    ledger_doc = test_files["ledger"]
    bank_doc = test_files["bank"]

    req = ChatRequest(
        message="What is the difference between the General Ledger and Bank Statement?",
        file_ids=[ledger_doc.id, bank_doc.id],
    )
    resp = asyncio.run(AssistantService.process_chat(db, req))

    msg = resp.message
    assert "50.0%" in msg or "50%" in msg
    assert "13,850" in msg
    assert "TXN1002" in msg
    assert "TXN1004" in msg or "12,000" in msg


def test_assistant_reconciliation_missing_bank(db: Session, test_files):
    ledger_doc = test_files["ledger"]
    bank_doc = test_files["bank"]

    req = ChatRequest(
        message="Show the transactions missing from the bank.",
        file_ids=[ledger_doc.id, bank_doc.id],
    )
    resp = asyncio.run(AssistantService.process_chat(db, req))

    msg = resp.message
    assert "TXN1004" in msg
    assert "TXN1006" in msg
    assert "15,200" in msg or ("12,000" in msg and "3,200" in msg)


def test_assistant_reconciliation_missing_ledger(db: Session, test_files):
    ledger_doc = test_files["ledger"]
    bank_doc = test_files["bank"]

    req = ChatRequest(
        message="Show the transactions missing from the ledger.",
        file_ids=[ledger_doc.id, bank_doc.id],
    )
    resp = asyncio.run(AssistantService.process_chat(db, req))

    msg = resp.message
    assert "TXN1007" in msg
    assert "450" in msg
    assert "Bank Service Charges" in msg or "service charge" in msg.lower()


def test_assistant_reconciliation_date_mismatches(db: Session, test_files):
    ledger_doc = test_files["ledger"]
    bank_doc = test_files["bank"]

    req = ChatRequest(
        message="Are there any date mismatches?",
        file_ids=[ledger_doc.id, bank_doc.id],
    )
    resp = asyncio.run(AssistantService.process_chat(db, req))

    msg = resp.message
    assert "No date mismatches" in msg or "0" in msg


def test_assistant_reconciliation_reference_mismatches(db: Session, test_files):
    ledger_doc = test_files["ledger"]
    bank_doc = test_files["bank"]

    req = ChatRequest(
        message="Are there any reference mismatches?",
        file_ids=[ledger_doc.id, bank_doc.id],
    )
    resp = asyncio.run(AssistantService.process_chat(db, req))

    msg = resp.message
    assert "No reference" in msg or "matching" in msg.lower()


def test_assistant_reconciliation_why_failing(db: Session, test_files):
    ledger_doc = test_files["ledger"]
    bank_doc = test_files["bank"]

    req = ChatRequest(
        message="Why is reconciliation failing?",
        file_ids=[ledger_doc.id, bank_doc.id],
    )
    resp = asyncio.run(AssistantService.process_chat(db, req))

    msg = resp.message
    assert "50.0%" in msg or "50%" in msg
    assert "13,850" in msg
    assert "TXN1002" in msg
    assert "TXN1004" in msg
    assert "TXN1007" in msg


def test_assistant_reconciliation_largest_mismatch(db: Session, test_files):
    ledger_doc = test_files["ledger"]
    bank_doc = test_files["bank"]

    req = ChatRequest(
        message="Where is the largest mismatch?",
        file_ids=[ledger_doc.id, bank_doc.id],
    )
    resp = asyncio.run(AssistantService.process_chat(db, req))

    msg = resp.message
    assert "TXN1004" in msg or "12,000" in msg


# =============================================================================
# 6. Negative Guard Tests
# =============================================================================

def test_assistant_missing_bank_guard(db: Session, test_files):
    ledger_doc = test_files["ledger"]

    req = ChatRequest(
        message="Find any amount mismatches between the General Ledger and Bank Statement.",
        file_ids=[ledger_doc.id],
    )
    resp = asyncio.run(AssistantService.process_chat(db, req))

    msg = resp.message
    assert "Bank Statement" in msg
    assert "LiteBooks_Reconciliation_Test_Data.xlsx" in msg or "General Ledger" in msg
    assert "attach" in msg.lower() or "upload" in msg.lower()


def test_assistant_missing_ledger_guard(db: Session, test_files):
    bank_doc = test_files["bank"]

    req = ChatRequest(
        message="Find any amount mismatches between the General Ledger and Bank Statement.",
        file_ids=[bank_doc.id],
    )
    resp = asyncio.run(AssistantService.process_chat(db, req))

    msg = resp.message
    assert "General Ledger" in msg
    assert "bankFile.xlsx" in msg or "Bank Statement" in msg
    assert "attach" in msg.lower() or "upload" in msg.lower()
