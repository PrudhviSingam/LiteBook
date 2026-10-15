import hashlib
import json
import logging
import re
import shutil
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.document import Document
from app.models.history import HistoryRecord
from app.models.reconciliation import CorrectionProposal, ReconciliationSession
from app.modules.reconciliation.correction import CorrectionProposalFactory
from app.modules.reconciliation.schemas import (
    CorrectionDecisionRequest,
    CorrectionProposalSchema,
    CorrectionStatusEnum,
    GenerateCorrectedWorkbookResponse,
)
from app.modules.reconciliation.workbook_inspector import (
    format_indian_currency,
    parse_indian_currency,
)
from app.services.history_service import HistoryService

logger = logging.getLogger(__name__)


class CorrectionService:
    """Orchestrates correction proposals, accountant approval lifecycle,

    non-destructive Excel copy generation, AI Correction Log worksheet creation,
    and Document/History integration.
    """

    @classmethod
    def list_or_create_proposals_for_session(
        cls, db: Session, session_id: int, user_id: Optional[int] = None
    ) -> List[CorrectionProposalSchema]:
        """Lists existing proposals for a reconciliation session or auto-generates

        initial candidates from discrepancies.
        """
        session = (
            db.query(ReconciliationSession)
            .filter(ReconciliationSession.id == session_id)
            .first()
        )
        if not session:
            return []
        if user_id is not None and session.user_id is not None and session.user_id != user_id:
            raise HTTPException(status_code=403, detail="Access denied. You do not own this reconciliation session.")

        existing = (
            db.query(CorrectionProposal)
            .filter(CorrectionProposal.session_id == session_id)
            .order_by(CorrectionProposal.created_at.asc())
            .all()
        )
        if existing:
            return [cls._model_to_schema(m) for m in existing]

        # Auto-seed proposals from discrepancies in session details
        data = json.loads(session.details_json)
        txns = data.get("transactions", [])
        proposals: List[CorrectionProposalSchema] = []

        counter = 1
        for t in txns:
            status = t.get("status")
            if status != "MATCHED":
                # Create a candidate proposal
                bank_amt = t.get("bank_amount") or 0.0
                ledger_amt = t.get("ledger_amount") or 0.0
                desc = t.get("description") or "Reconciliation Item"
                ref = t.get("reference") or f"TXN-{counter:03d}"
                diff = t.get("difference") or round(bank_amt - ledger_amt, 2)

                orig_val = ledger_amt if ledger_amt > 0 else bank_amt
                prop_val = bank_amt if ledger_amt > 0 else ledger_amt

                cid = f"COR-{session_id:04d}-{counter:03d}"
                prop = CorrectionProposalFactory.create(
                    session_id=session_id,
                    sheet_name="Balance Sheet" if "payables" in desc.lower() or "lease" in desc.lower() else "Ledger",
                    row_reference=counter + 10,
                    field="Amount",
                    original_value=orig_val,
                    proposed_value=prop_val,
                    reason_category=status,
                    explanation=(
                        f"The system detected a discrepancy of {format_indian_currency(diff)} for '{desc}'. "
                        "This may be caused by an omitted transaction, entry error, or timing difference. "
                        "Please verify before approving."
                    ),
                    evidence=[
                        f"Matched description / reference: '{desc}'",
                        f"Recorded value: {format_indian_currency(orig_val)}",
                        f"Expected comparison value: {format_indian_currency(prop_val)}",
                        f"Variance: {format_indian_currency(diff)}",
                    ],
                    recommended_verification=[
                        f"Verify supporting voucher/statement for '{desc}'",
                        "Confirm if tax deduction, bank charge, or debit note was issued",
                    ],
                    confidence="HIGH" if abs(diff) > 0 else "MEDIUM",
                    document_id=session.ledger_document_id or session.bank_document_id,
                    correction_id=cid,
                )

                # Persist in DB
                db_model = CorrectionProposal(
                    user_id=session.user_id or user_id,
                    correction_id=prop.correction_id,
                    session_id=prop.session_id,
                    document_id=prop.document_id,
                    sheet_name=prop.sheet_name,
                    row_reference=str(prop.row_reference),
                    field=prop.field,
                    original_value=str(prop.original_value),
                    proposed_value=str(prop.proposed_value),
                    difference=prop.difference,
                    reason_category=prop.reason_category,
                    explanation=prop.explanation,
                    evidence_json=json.dumps(prop.evidence),
                    recommended_verification_json=json.dumps(prop.recommended_verification),
                    confidence=prop.confidence,
                    status=prop.status,
                )
                db.add(db_model)
                proposals.append(prop)
                counter += 1

        if proposals:
            db.commit()
            # Log history
            HistoryService.record_activity(
                db=db,
                user_id=session.user_id or user_id,
                module="reconciliation",
                action="discrepancy_detected",
                target_name=f"Session {session_id} Candidates",
                status="completed",
                result_summary=f"Proposed {len(proposals)} corrections for review",
                user_action="Automatic Proposal Generation",
                report_reference=f"REC-{session_id:04d}",
            )

        return proposals

    @classmethod
    def update_proposal_decision(
        cls,
        db: Session,
        correction_id: str,
        decision_req: CorrectionDecisionRequest,
        user_id: Optional[int] = None,
    ) -> Optional[CorrectionProposalSchema]:
        """Updates the accountant decision on a correction proposal (CONFIRMED, REJECTED, MANUAL_REVIEW)."""
        prop = (
            db.query(CorrectionProposal)
            .filter(CorrectionProposal.correction_id == correction_id)
            .first()
        )
        if not prop:
            return None

        if user_id is not None and prop.user_id is not None and prop.user_id != user_id:
            raise HTTPException(status_code=403, detail="Access denied. You do not own this correction proposal.")

        dec = decision_req.decision.upper()
        if dec == "CONFIRMED":
            prop.status = CorrectionStatusEnum.CONFIRMED.value
            prop.accountant_decision = "approved"
            hist_action = "correction_approved"
            summary_msg = f"Correction {correction_id} approved by accountant."
        elif dec == "REJECTED":
            prop.status = CorrectionStatusEnum.REJECTED.value
            prop.accountant_decision = "rejected"
            hist_action = "correction_rejected"
            summary_msg = f"Correction {correction_id} rejected by accountant."
        else:
            prop.status = CorrectionStatusEnum.PROPOSED.value
            prop.accountant_decision = "review_manually"
            hist_action = "correction_review_deferred"
            summary_msg = f"Correction {correction_id} flagged for manual review."

        prop.decision_notes = decision_req.notes
        db.commit()
        db.refresh(prop)

        # Log history
        HistoryService.record_activity(
            db=db,
            user_id=user_id or prop.user_id,
            module="reconciliation",
            action=hist_action,
            target_name=f"{correction_id} ({prop.sheet_name})",
            status="completed",
            result_summary=summary_msg,
            user_action="Accountant decision",
            report_reference=f"REC-{prop.session_id:04d}" if prop.session_id else None,
        )

        return cls._model_to_schema(prop)

    @classmethod
    def generate_corrected_workbook(
        cls,
        db: Session,
        session_id: int,
        accountant_name: str = "Partner Accountant",
        user_id: Optional[int] = None,
    ) -> GenerateCorrectedWorkbookResponse:
        """Generates a timestamped, non-destructive copy of the original Excel workbook,

        applying explicitly approved corrections and appending an 'AI Correction Log' worksheet.
        """
        session = (
            db.query(ReconciliationSession)
            .filter(ReconciliationSession.id == session_id)
            .first()
        )
        if not session:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Reconciliation session {session_id} not found.")

        if user_id is not None and session.user_id is not None and session.user_id != user_id:
            raise HTTPException(status_code=403, detail="Access denied. You do not own this reconciliation session.")

        # 1. Fetch confirmed/approved corrections (Strictly requires explicit CONFIRMED status)
        confirmed_props = (
            db.query(CorrectionProposal)
            .filter(
                CorrectionProposal.session_id == session_id,
                CorrectionProposal.status == CorrectionStatusEnum.CONFIRMED.value,
            )
            .all()
        )

        if not confirmed_props:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail="Correction must be explicitly approved by the accountant before a corrected file can be generated.",
            )

        # 2. Locate original workbook
        orig_doc = None
        orig_path: Optional[Path] = None
        if session.ledger_document_id:
            orig_doc = db.query(Document).filter(Document.id == session.ledger_document_id).first()
        if not orig_doc and session.bank_document_id:
            orig_doc = db.query(Document).filter(Document.id == session.bank_document_id).first()

        if orig_doc and Path(orig_doc.file_path).exists():
            orig_path = Path(orig_doc.file_path)
            orig_filename = orig_doc.original_filename
        else:
            orig_filename = session.ledger_file_name or session.bank_statement_name or "Financial_Workbook.xlsx"

        # Record original file SHA-256 for immutability verification
        orig_sha256_before = None
        if orig_path and orig_path.exists():
            with open(orig_path, "rb") as f:
                orig_sha256_before = hashlib.sha256(f.read()).hexdigest()

        # Timestamped generated filename
        timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        base_name = Path(orig_filename).stem
        # Clean any existing "Corrected" suffixes to prevent compounding
        base_name = re.sub(r"_Corrected_\d+_\d+", "", base_name)
        generated_filename = f"{base_name}_Corrected_{timestamp_str}.xlsx"

        output_dir = settings.uploads_dir
        dest_path = output_dir / f"{uuid.uuid4().hex}_{generated_filename}"

        # 3. Load or build workbook preserving original structure
        if orig_path and orig_path.exists() and orig_path.suffix.lower() in (".xlsx", ".xlsm"):
            # Load existing workbook with styles preserved
            wb = openpyxl.load_workbook(orig_path, data_only=False)
        else:
            # Build clean workbook structure matching the session data
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = confirmed_props[0].sheet_name or "Balance Sheet"
            # Seed header and sample rows
            ws.append(["Account / Description", "Reference", "Amount"])
            for p in confirmed_props:
                ws.append([f"Account Item {p.row_reference}", f"REF-{p.row_reference}", float(p.original_value or 0)])

        # 4. Apply confirmed corrections strictly to target cells with cell verification
        for p in confirmed_props:
            sheet_target = p.sheet_name
            if sheet_target not in wb.sheetnames:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail=f"Target sheet '{sheet_target}' not found in workbook.",
                )
            ws = wb[sheet_target]

            # Locate target cell: either by exact cell_address or row/column match
            target_cell = None
            if p.cell_address and p.cell_address in ws:
                target_cell = ws[p.cell_address]
            else:
                # Find cell by row_reference
                try:
                    row_idx = int(p.row_reference)
                except ValueError:
                    row_idx = 2

                # Determine column index: search header row for field or use column 2/3
                col_idx = 2
                field_name = str(p.field).lower()
                for col in range(1, ws.max_column + 1):
                    cval = str(ws.cell(row=1, column=col).value or "").lower()
                    if field_name in cval or ("amount" in cval and field_name == "amount"):
                        col_idx = col
                        break
                target_cell = ws.cell(row=row_idx, column=col_idx)

                # If default row/col does not match original_value, search sheet for exact cell
                current_cell_val = parse_indian_currency(target_cell.value)
                expected_orig_val = parse_indian_currency(p.original_value)
                if current_cell_val is None or (expected_orig_val is not None and abs(current_cell_val - expected_orig_val) > 0.01):
                    for r in range(1, min(ws.max_row + 1, 100)):
                        for c in range(1, min(ws.max_column + 1, 20)):
                            val = parse_indian_currency(ws.cell(row=r, column=c).value)
                            if val is not None and expected_orig_val is not None and abs(val - expected_orig_val) < 0.01:
                                target_cell = ws.cell(row=r, column=c)
                                p.cell_address = f"{openpyxl.utils.get_column_letter(c)}{r}"
                                break
                        if p.cell_address:
                            break

            # Integrity check: read current value and verify it matches stored original_value
            current_cell_val = parse_indian_currency(target_cell.value)
            expected_orig_val = parse_indian_currency(p.original_value)
            if current_cell_val is not None and expected_orig_val is not None:
                if abs(current_cell_val - expected_orig_val) > 0.01:
                    raise HTTPException(
                        status_code=status.HTTP_409_CONFLICT,
                        detail="The workbook changed after the correction proposal was created. The proposed correction was not applied.",
                    )

            # Update value safely
            new_val = parse_indian_currency(p.proposed_value)
            target_cell.value = new_val if new_val is not None else p.proposed_value
            # Mark cell with a subtle accounting audit fill (soft emerald)
            target_cell.fill = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")


        # 5. Append or update 'AI Correction Log' worksheet
        log_sheet_name = "AI Correction Log"
        if log_sheet_name in wb.sheetnames:
            del wb[log_sheet_name]
        ws_log = wb.create_sheet(title=log_sheet_name)


        headers = [
            "Correction ID",
            "Date/Time",
            "Sheet",
            "Row",
            "Field",
            "Original Value",
            "New Value",
            "Difference",
            "Reason Category",
            "AI Explanation",
            "Evidence Summary",
            "Confidence",
            "Approved By",
            "Status",
        ]
        ws_log.append(headers)

        # Style header row
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        for col_num in range(1, len(headers) + 1):
            cell = ws_log.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Insert log rows
        now_iso = datetime.now().strftime("%d-%m-%Y %H:%M:%S")
        for p in confirmed_props:
            evidence_str = "; ".join(json.loads(p.evidence_json or "[]"))
            ws_log.append([
                p.correction_id,
                now_iso,
                p.sheet_name,
                p.row_reference,
                p.field,
                float(p.original_value or 0),
                float(p.proposed_value or 0),
                float(p.difference or 0),
                p.reason_category,
                p.explanation,
                evidence_str,
                p.confidence,
                accountant_name,
                "GENERATED",
            ])

        # Auto-fit column widths
        for col in ws_log.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws_log.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # 6. Save new file without touching original
        wb.save(dest_path)
        file_size = dest_path.stat().st_size

        # Compute checksum
        sha256 = hashlib.sha256()
        with open(dest_path, "rb") as f:
            for chunk in iter(lambda: f.read(4096), b""):
                sha256.update(chunk)
        checksum = sha256.hexdigest()

        # Verify original file remains strictly immutable
        if orig_path and orig_path.exists() and orig_sha256_before:
            with open(orig_path, "rb") as f:
                orig_sha256_after = hashlib.sha256(f.read()).hexdigest()
            if orig_sha256_before != orig_sha256_after:
                if dest_path.exists():
                    dest_path.unlink()
                raise HTTPException(
                    status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                    detail="Immutability safety violation: original file was modified during copy generation.",
                )

        # Update confirmed proposals to GENERATED status
        for p in confirmed_props:
            p.status = CorrectionStatusEnum.GENERATED.value
            p.original_sha256 = orig_sha256_before
            p.corrected_sha256 = checksum
            p.reviewed_at = datetime.now(timezone.utc)

        user_id = getattr(session, "user_id", None)

        # 7. Register generated file in Documents module
        doc_record = Document(
            user_id=user_id,
            original_filename=generated_filename,
            stored_filename=dest_path.name,
            file_path=str(dest_path),
            file_size=file_size,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            category="financial",
            status="ready",
            checksum=checksum,
            original_sha256=orig_sha256_before,
            is_corrected_copy=True,
            notes=json.dumps({
                "source": "reconciliation_correction",
                "session_id": session_id,
                "original_file": orig_filename,
                "corrections_count": len(confirmed_props),
                "generated_at": datetime.now(timezone.utc).isoformat(),
            }),
        )
        db.add(doc_record)
        db.commit()
        db.refresh(doc_record)


        # 8. Log in History
        HistoryService.record_activity(
            db=db,
            module="reconciliation",
            action="corrected_file_generated",
            target_name=generated_filename,
            status="completed",
            result_summary=(
                f"Generated corrected copy '{generated_filename}' ({round(file_size / 1024, 1)} KB) "
                f"with {len(confirmed_props)} approved correction(s) and AI Correction Log."
            ),
            user_action="Approve & Generate Copy",
            report_reference=f"REC-{session_id:04d}",
            document_id=doc_record.id,
        )

        return GenerateCorrectedWorkbookResponse(
            success=True,
            original_filename=orig_filename,
            generated_filename=generated_filename,
            download_url=f"/api/v1/documents/{doc_record.id}/download",
            document_id=doc_record.id,
            applied_corrections_count=len(confirmed_props),
            corrections_sheet_name=log_sheet_name,
            message=(
                f"Successfully generated non-destructive copy '{generated_filename}' "
                f"with {len(confirmed_props)} approved correction(s) and AI Correction Log sheet."
            ),
        )

    @staticmethod
    def _model_to_schema(m: CorrectionProposal) -> CorrectionProposalSchema:
        return CorrectionProposalSchema(
            correction_id=m.correction_id,
            session_id=m.session_id,
            document_id=m.document_id,
            sheet_name=m.sheet_name,
            cell_address=m.cell_address,
            row_reference=m.row_reference,

            field=m.field,
            original_value=float(m.original_value) if m.original_value else 0.0,
            proposed_value=float(m.proposed_value) if m.proposed_value else 0.0,
            difference=m.difference,
            reason_category=m.reason_category,
            explanation=m.explanation,
            evidence=json.loads(m.evidence_json or "[]"),
            recommended_verification=json.loads(m.recommended_verification_json or "[]"),
            confidence=m.confidence,
            status=m.status,
            accountant_decision=m.accountant_decision,
            decision_notes=m.decision_notes,
            created_at=m.created_at.isoformat() if m.created_at else None,
        )
