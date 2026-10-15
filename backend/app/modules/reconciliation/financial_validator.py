import re
import uuid
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.utils import get_column_letter

from app.modules.reconciliation.schemas import (
    DiscrepancySeverityEnum,
    DiscrepancyTypeEnum,
    DocumentTypeEnum,
    FinancialCellProvenance,
    FinancialRow,
    FinancialTable,
    StandardDiscrepancy,
)


def format_currency_inr(amount: Optional[Any]) -> str:
    """Formats a float or numeric string into standard Indian numbering format (e.g. ₹20,14,02,126.89)."""
    if amount is None:
        return "₹0.00"
    if isinstance(amount, str):
        amount_str = amount.strip()
        if amount_str.startswith("₹") or amount_str.startswith("Rs"):
            return amount_str
        try:
            amount = float(amount_str.replace(",", ""))
        except Exception:
            return amount_str

    try:
        is_neg = amount < 0
        abs_amt = abs(amount)
        parts = f"{abs_amt:.2f}".split(".")
        integer_part = parts[0]
        decimal_part = parts[1]
    except Exception:
        return str(amount)

    if len(integer_part) <= 3:
        formatted_int = integer_part
    else:
        last_three = integer_part[-3:]
        remaining = integer_part[:-3]
        groups = []
        while len(remaining) > 2:
            groups.insert(0, remaining[-2:])
            remaining = remaining[:-2]
        if remaining:
            groups.insert(0, remaining)
        formatted_int = ",".join(groups) + "," + last_three

    res = f"₹{formatted_int}.{decimal_part}"
    return f"-{res}" if is_neg else res


class FinancialValidator:
    """Deterministic financial validation engine.

    Never allows an AI or LLM to calculate authoritative totals or variances.
    Executes exact mathematical validation rules across financial statements.
    """

    @classmethod
    def validate_workbook(
        cls,
        wb_data: openpyxl.Workbook,
        wb_formula: Optional[openpyxl.Workbook] = None,
        tables: Optional[List[FinancialTable]] = None,
    ) -> List[StandardDiscrepancy]:
        """Runs all deterministic validation passes over the openpyxl workbook."""
        discrepancies: List[StandardDiscrepancy] = []

        # 1. Balance Sheet validation
        bs_discrepancies = cls.validate_balance_sheets(wb_data, wb_formula)
        discrepancies.extend(bs_discrepancies)

        # 2. Trial Balance validation
        tb_discrepancies = cls.validate_trial_balance(wb_data, wb_formula)
        discrepancies.extend(tb_discrepancies)

        # 3. Profit & Loss validation
        pl_discrepancies = cls.validate_profit_loss(wb_data, wb_formula)
        discrepancies.extend(pl_discrepancies)

        # 4. Formula awareness checks
        if wb_formula:
            formula_discrepancies = cls.validate_formulas(wb_data, wb_formula)
            discrepancies.extend(formula_discrepancies)

        return discrepancies

    # -------------------------------------------------------------------------
    # 1. BALANCE SHEET VALIDATION
    # -------------------------------------------------------------------------

    @classmethod
    def validate_balance_sheets(
        cls,
        wb_data: openpyxl.Workbook,
        wb_formula: Optional[openpyxl.Workbook] = None,
    ) -> List[StandardDiscrepancy]:
        """Validates that Total Assets == Total Liabilities + Equity.

        Supports standard Schedule III Indian format and horizontal formats.
        """
        findings: List[StandardDiscrepancy] = []

        for sheet_name in wb_data.sheetnames:
            sheet_lower = sheet_name.lower()
            ws_data = wb_data[sheet_name]

            # Check if this sheet is likely a Balance Sheet
            # Inspect first 25 rows for balance sheet markers
            sample_text = ""
            for r in range(1, min(ws_data.max_row + 1, 25)):
                row_vals = [str(ws_data.cell(r, c).value or "") for c in range(1, min(ws_data.max_column + 1, 10))]
                sample_text += " " + " ".join(row_vals).lower()

            is_bs = "balance sheet" in sheet_lower or "statement of assets and liabilities" in sample_text or (
                "total assets" in sample_text and ("total liabilities" in sample_text or "equity and liabilities" in sample_text)
            )

            if not is_bs:
                continue

            # Locate Total Assets, Total Liabilities, and periods
            assets_info = cls._find_row_total(ws_data, keywords=["total assets", "total asset"])
            liab_info = cls._find_row_total(
                ws_data,
                keywords=[
                    "total liabilities",
                    "total equity and liabilities",
                    "total equity & liabilities",
                    "total liabilities and equity",
                ],
            )

            if assets_info and liab_info:
                # Compare for matching columns
                for col_idx in assets_info["values"]:
                    if col_idx in liab_info["values"]:
                        assets_val = assets_info["values"][col_idx]
                        liab_val = liab_info["values"][col_idx]
                        asset_cell = f"{get_column_letter(col_idx)}{assets_info['row']}"
                        liab_cell = f"{get_column_letter(col_idx)}{liab_info['row']}"

                        variance = round(abs(liab_val - assets_val), 2)
                        if variance > 0.01:
                            col_hdr = cls._get_column_header(ws_data, col_idx)
                            period_context = f" for {col_hdr}" if col_hdr else ""

                            if liab_val > assets_val:
                                desc = f"Liabilities and equity exceed total assets by {format_currency_inr(variance)}{period_context}."
                            else:
                                desc = f"Total assets exceed liabilities and equity by {format_currency_inr(variance)}{period_context}."

                            findings.append(
                                StandardDiscrepancy(
                                    id=f"disc_bs_{uuid.uuid4().hex[:8]}",
                                    type=DiscrepancyTypeEnum.BALANCE_SHEET_IMBALANCE,
                                    severity=DiscrepancySeverityEnum.CRITICAL,
                                    status="DETECTED",
                                    description=desc,
                                    expected_value=format_currency_inr(liab_val),
                                    actual_value=format_currency_inr(assets_val),
                                    difference=variance,
                                    currency="INR",
                                    sheet_name=sheet_name,
                                    cell_address=asset_cell,
                                    source_reference=f"Assets: {sheet_name}!{asset_cell} vs Liabilities: {sheet_name}!{liab_cell}",
                                    confidence="HIGH",
                                    evidence=[
                                        f"Stated Total Assets: {format_currency_inr(assets_val)} at {sheet_name}!{asset_cell}",
                                        f"Stated Total Liabilities + Equity: {format_currency_inr(liab_val)} at {sheet_name}!{liab_cell}",
                                        f"Calculated variance = {format_currency_inr(variance)}",
                                    ],
                                    possible_causes=[
                                        "Possible cause: missing line item in asset schedule.",
                                        "Possible cause: duplicated line item in liabilities or equity.",
                                        "Possible cause: incorrect total formula or range definition.",
                                        "Possible cause: rounding difference across sub-schedules.",
                                        "Possible cause: manually entered total differing from line items.",
                                        "Possible cause: classification issue between capital, reserves, or current liabilities.",
                                    ],
                                    recommended_actions=[
                                        f"Review asset schedule items above {asset_cell}.",
                                        f"Verify liability schedule breakdown above {liab_cell}.",
                                        "Check supporting ledger accounts for trade payables, fixed assets, and bank balances.",
                                    ],
                                )
                            )

        return findings

    # -------------------------------------------------------------------------
    # 2. TRIAL BALANCE VALIDATION
    # -------------------------------------------------------------------------

    @classmethod
    def validate_trial_balance(
        cls,
        wb_data: openpyxl.Workbook,
        wb_formula: Optional[openpyxl.Workbook] = None,
    ) -> List[StandardDiscrepancy]:
        """Validates that Total Debits == Total Credits in Trial Balance sheets."""
        findings: List[StandardDiscrepancy] = []

        for sheet_name in wb_data.sheetnames:
            ws_data = wb_data[sheet_name]
            sheet_lower = sheet_name.lower()

            sample_text = ""
            for r in range(1, min(ws_data.max_row + 1, 20)):
                row_vals = [str(ws_data.cell(r, c).value or "") for c in range(1, min(ws_data.max_column + 1, 10))]
                sample_text += " " + " ".join(row_vals).lower()

            is_tb = "trial balance" in sheet_lower or "trial balance" in sample_text or (
                "debit" in sample_text and "credit" in sample_text and ("total" in sample_text or "closing balance" in sample_text)
            )

            if not is_tb:
                continue

            # Locate Total row
            for r in range(ws_data.max_row, 1, -1):
                row_vals = [str(ws_data.cell(r, c).value or "").strip().lower() for c in range(1, ws_data.max_column + 1)]
                first_few = " ".join(row_vals[:3])
                if "total" in first_few or "grand total" in first_few:
                    num_vals: List[Tuple[int, float]] = []
                    for c in range(1, ws_data.max_column + 1):
                        v = ws_data.cell(r, c).value
                        parsed = cls._parse_numeric(v)
                        if parsed is not None and parsed > 0:
                            num_vals.append((c, parsed))

                    if len(num_vals) >= 2:
                        col_a, val_a = num_vals[0]
                        col_b, val_b = num_vals[1]
                        diff = round(abs(val_a - val_b), 2)
                        cell_a = f"{get_column_letter(col_a)}{r}"
                        cell_b = f"{get_column_letter(col_b)}{r}"

                        if diff > 0.01:
                            findings.append(
                                StandardDiscrepancy(
                                    id=f"disc_tb_{uuid.uuid4().hex[:8]}",
                                    type=DiscrepancyTypeEnum.TRIAL_BALANCE_IMBALANCE,
                                    severity=DiscrepancySeverityEnum.CRITICAL,
                                    status="DETECTED",
                                    description=f"Trial balance is out of balance by {format_currency_inr(diff)}.",
                                    expected_value=format_currency_inr(val_a),
                                    actual_value=format_currency_inr(val_b),
                                    difference=diff,
                                    currency="INR",
                                    sheet_name=sheet_name,
                                    cell_address=cell_b,
                                    source_reference=f"{sheet_name}!{cell_a} vs {sheet_name}!{cell_b}",
                                    confidence="HIGH",
                                    evidence=[
                                        f"Debit Total: {format_currency_inr(val_a)} ({sheet_name}!{cell_a})",
                                        f"Credit Total: {format_currency_inr(val_b)} ({sheet_name}!{cell_b})",
                                        f"Calculated variance: {format_currency_inr(diff)}",
                                    ],
                                    possible_causes=[
                                        "Possible cause: single-sided journal entry posted without matching credit/debit.",
                                        "Possible cause: incorrect ledger closing balance extraction.",
                                        "Possible cause: transposition error or omitted ledger account.",
                                    ],
                                    recommended_actions=[
                                        "Inspect journal entries for the period for unbalanced postings.",
                                        "Verify opening balance sheet carry-forwards.",
                                    ],
                                )
                            )
                    break

        return findings

    # -------------------------------------------------------------------------
    # 3. PROFIT & LOSS VALIDATION
    # -------------------------------------------------------------------------

    @classmethod
    def validate_profit_loss(
        cls,
        wb_data: openpyxl.Workbook,
        wb_formula: Optional[openpyxl.Workbook] = None,
    ) -> List[StandardDiscrepancy]:
        """Validates Net Profit = Total Revenue - Total Expenses where structured."""
        findings: List[StandardDiscrepancy] = []

        for sheet_name in wb_data.sheetnames:
            ws_data = wb_data[sheet_name]
            sheet_lower = sheet_name.lower()

            sample_text = ""
            for r in range(1, min(ws_data.max_row + 1, 20)):
                row_vals = [str(ws_data.cell(r, c).value or "") for c in range(1, min(ws_data.max_column + 1, 10))]
                sample_text += " " + " ".join(row_vals).lower()

            is_pl = "profit" in sheet_lower or "loss" in sheet_lower or "p&l" in sheet_lower or (
                "revenue from operations" in sample_text and "expenses" in sample_text
            )

            if not is_pl:
                continue

            rev_info = cls._find_row_total(ws_data, ["revenue from operations", "total revenue", "revenue"])
            other_inc_info = cls._find_row_total(ws_data, ["other income", "other revenue"])
            exp_info = cls._find_row_total(ws_data, ["total expenses", "total expense"])
            profit_info = cls._find_row_total(ws_data, ["net profit", "net profit before tax", "profit before tax", "profit for the year"])

            if rev_info and exp_info and profit_info:
                for col_idx in rev_info["values"]:
                    if col_idx in exp_info["values"] and col_idx in profit_info["values"]:
                        rev = rev_info["values"][col_idx]
                        other_inc = other_inc_info["values"].get(col_idx, 0.0) if other_inc_info else 0.0
                        total_rev = rev + other_inc
                        exp = exp_info["values"][col_idx]
                        stated_profit = profit_info["values"][col_idx]
                        expected_profit = round(total_rev - exp, 2)
                        diff = round(abs(stated_profit - expected_profit), 2)

                        if diff > 0.01:
                            cell_addr = f"{get_column_letter(col_idx)}{profit_info['row']}"
                            findings.append(
                                StandardDiscrepancy(
                                    id=f"disc_pl_{uuid.uuid4().hex[:8]}",
                                    type=DiscrepancyTypeEnum.AMOUNT_MISMATCH,
                                    severity=DiscrepancySeverityEnum.HIGH,
                                    status="DETECTED",
                                    description=f"Net Profit calculation mismatch of {format_currency_inr(diff)}.",
                                    expected_value=format_currency_inr(expected_profit),
                                    actual_value=format_currency_inr(stated_profit),
                                    difference=diff,
                                    currency="INR",
                                    sheet_name=sheet_name,
                                    cell_address=cell_addr,
                                    confidence="HIGH",
                                    evidence=[
                                        f"Revenue: {format_currency_inr(rev)}",
                                        f"Expenses: {format_currency_inr(exp)}",
                                        f"Calculated Net Profit: {format_currency_inr(expected_profit)}",
                                        f"Stated Net Profit: {format_currency_inr(stated_profit)} at {sheet_name}!{cell_addr}",
                                    ],
                                    possible_causes=[
                                        "Possible cause: other income omitted from revenue calculation.",
                                        "Possible cause: tax provision or depreciation deducted without distinct row.",
                                        "Possible cause: arithmetic error in Profit & Loss compilation.",
                                    ],
                                    recommended_actions=[
                                        "Check if Other Income was added to Revenue from Operations.",
                                        "Confirm whether exceptional items or taxes were subtracted.",
                                    ],
                                )
                            )

        return findings

    # -------------------------------------------------------------------------
    # 4. FORMULA AWARENESS
    # -------------------------------------------------------------------------

    @classmethod
    def validate_formulas(
        cls,
        wb_data: openpyxl.Workbook,
        wb_formula: openpyxl.Workbook,
    ) -> List[StandardDiscrepancy]:
        """Detects hardcoded totals, missing formulas, and formula range discrepancies."""
        findings: List[StandardDiscrepancy] = []

        for sheet_name in wb_data.sheetnames:
            if sheet_name not in wb_formula.sheetnames:
                continue
            ws_data = wb_data[sheet_name]
            ws_form = wb_formula[sheet_name]

            for r in range(1, ws_data.max_row + 1):
                row_label = str(ws_data.cell(r, 1).value or ws_data.cell(r, 2).value or "").strip().lower()
                is_total_row = "total" in row_label or "subtotal" in row_label or "grand total" in row_label

                if not is_total_row:
                    continue

                for c in range(2, min(ws_data.max_column + 1, 10)):
                    val = ws_data.cell(r, c).value
                    formula_val = str(ws_form.cell(r, c).value or "")

                    num = cls._parse_numeric(val)
                    if num is not None and abs(num) > 0:
                        cell_addr = f"{get_column_letter(c)}{r}"

                        # Check if total is hardcoded (no leading '=')
                        if not formula_val.startswith("="):
                            has_items_above = False
                            for prev_r in range(max(1, r - 5), r):
                                prev_num = cls._parse_numeric(ws_data.cell(prev_r, c).value)
                                if prev_num is not None and prev_num > 0:
                                    has_items_above = True
                                    break

                            if has_items_above:
                                findings.append(
                                    StandardDiscrepancy(
                                        id=f"disc_form_{uuid.uuid4().hex[:8]}",
                                        type=DiscrepancyTypeEnum.FORMULA_ERROR,
                                        severity=DiscrepancySeverityEnum.INFO,
                                        status="DETECTED",
                                        description=f"Total in {row_label.title()} appears to be hardcoded rather than dynamically calculated.",
                                        expected_value="Dynamic Excel formula (e.g., =SUM(...))",
                                        actual_value=format_currency_inr(num),
                                        difference=0.0,
                                        currency="INR",
                                        sheet_name=sheet_name,
                                        cell_address=cell_addr,
                                        confidence="MEDIUM",
                                        evidence=[
                                            f"Cell {sheet_name}!{cell_addr} contains static value {format_currency_inr(num)} instead of a SUM formula.",
                                        ],
                                        possible_causes=[
                                            "Possible cause: manually entered total rather than formula-driven summation.",
                                            "Possible cause: paste-as-values from external accounting export.",
                                        ],
                                        recommended_actions=[
                                            f"Verify if {sheet_name}!{cell_addr} equals the sum of line items above it.",
                                        ],
                                    )
                                )
                        else:
                            m = re.match(r"=SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", formula_val, re.IGNORECASE)
                            if m:
                                f_col_start, f_row_start, f_col_end, f_row_end = m.groups()
                                f_end = int(f_row_end)
                                if f_end < (r - 1):
                                    skipped_rows = (r - 1) - f_end
                                    findings.append(
                                        StandardDiscrepancy(
                                            id=f"disc_range_{uuid.uuid4().hex[:8]}",
                                            type=DiscrepancyTypeEnum.FORMULA_ERROR,
                                            severity=DiscrepancySeverityEnum.HIGH,
                                            status="DETECTED",
                                            description=f"Formula at {cell_addr} omits {skipped_rows} row(s) directly preceding the total.",
                                            expected_value=f"=SUM({f_col_start}{f_row_start}:{f_col_end}{r - 1})",
                                            actual_value=formula_val,
                                            difference=0.0,
                                            currency="INR",
                                            sheet_name=sheet_name,
                                            cell_address=cell_addr,
                                            confidence="HIGH",
                                            evidence=[
                                                f"Formula {formula_val} stops at row {f_end}, but total is at row {r}.",
                                            ],
                                            possible_causes=[
                                                "Possible cause: new line items inserted without updating the SUM formula range.",
                                            ],
                                            recommended_actions=[
                                                f"Update formula range in {cell_addr} to include omitted line items.",
                                            ],
                                        )
                                    )

        return findings

    # -------------------------------------------------------------------------
    # HELPER UTILITIES
    # -------------------------------------------------------------------------

    @classmethod
    def _parse_numeric(cls, val: Any) -> Optional[float]:
        if val is None:
            return None
        if isinstance(val, (int, float)):
            return float(val)
        s = str(val).strip()
        if not s:
            return None
        is_neg = False
        if s.startswith("(") and s.endswith(")"):
            is_neg = True
            s = s[1:-1].strip()
        elif s.startswith("-"):
            is_neg = True
            s = s[1:].strip()
        s = re.sub(r"[₹$,\s]", "", s)
        try:
            v = float(s)
            return -v if is_neg else v
        except ValueError:
            return None

    @classmethod
    def _find_row_total(cls, ws: Any, keywords: List[str]) -> Optional[Dict[str, Any]]:
        for r in range(1, ws.max_row + 1):
            row_vals = [str(ws.cell(r, c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]
            row_text = " ".join(row_vals)
            if any(k in row_text for k in keywords):
                values: Dict[int, float] = {}
                for c in range(1, ws.max_column + 1):
                    parsed = cls._parse_numeric(ws.cell(r, c).value)
                    if parsed is not None and abs(parsed) > 0:
                        values[c] = parsed
                if values:
                    return {"row": r, "values": values}
        return None

    @classmethod
    def _get_column_header(cls, ws: Any, col_idx: int) -> Optional[str]:
        for r in range(1, min(ws.max_row + 1, 10)):
            v = ws.cell(r, col_idx).value
            if v and str(v).strip():
                return str(v).strip()
        return None
